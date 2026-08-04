from __future__ import annotations

import sys
import tempfile
import unittest
import base64
import json
from contextlib import redirect_stdout
from io import BytesIO, StringIO
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from vipshop_hot_strategy_agent.analysis import analyze_workbook, write_analysis
from vipshop_hot_strategy_agent.cli import main as cli_main
from vipshop_hot_strategy_agent.collect import _expand_download_attachments, build_params
from vipshop_hot_strategy_agent.excel_export import export_workbook


SAMPLE_ROWS = [
    {"__sheet_name": "执行摘要", "报表来源": "test", "执行结果": "ok", "备注": "sample"},
    {
        "__sheet_name": "魔方罗盘销售明细",
        "款号": "A100",
        "货号": "A100-RED",
        "商品ID": "10001",
        "三级品类": "童装/T恤",
        "销售额": 12000,
        "销售数量": 120,
        "客户数": 80,
        "商详UV": 1000,
        "曝光人数": 5000,
        "点击人数": 600,
        "加购数": 150,
        "下单客户数": 90,
        "在售库存": 20,
        "库存天数": 5,
    },
    {
        "__sheet_name": "魔方罗盘销售明细",
        "款号": "B200",
        "货号": "B200-BLU",
        "商品ID": "10002",
        "三级品类": "童装/裤子",
        "销售额": 8000,
        "销售数量": 60,
        "客户数": 50,
        "商详UV": 700,
        "曝光人数": 3000,
        "点击人数": 260,
        "加购数": 90,
        "下单客户数": 55,
        "在售库存": 100,
        "库存天数": 30,
    },
    {
        "__sheet_name": "唯直达投放效果",
        "统计日期": "2026-08-02",
        "花费": 100,
        "曝光量": 10000,
        "点击量": 300,
        "24小时订单额(商家)": 900,
        "24小时销售额(商家)": 850,
        "24小时加购数(商家)": 60,
    },
    {
        "__sheet_name": "T-max效果",
        "商品ID": "10001",
        "曝光人数": 5000,
        "点击人数": 200,
        "销售额": 1800,
        "活动成本": 150,
        "ROI": 12,
    },
    {
        "__sheet_name": "中台礼金",
        "活动ID": "G1",
        "活动名称": "老友礼金",
        "运营方式": "礼金",
        "活动状态": "进行中",
        "触达用户数": 500,
        "已使用人数": 25,
        "使用量": 35,
        "销售额": 1600,
        "活动成本": 100,
        "订单数": 30,
        "客户数": 25,
        "计划开始": "2026-08-01 00:00:00",
        "计划结束": "2026-08-03 23:59:59",
    },
    {
        "__sheet_name": "中台购物车跨品类券",
        "活动ID": "C1",
        "活动名称": "购物车挽回",
        "运营方式": "跨品类券",
        "活动状态": "进行中",
        "触达用户数": 260,
        "已使用人数": 10,
        "用券量": 15,
        "销售额": 700,
        "活动成本": 140,
        "订单数": 12,
        "客户数": 10,
        "计划开始": "2026-08-01 00:00:00",
        "计划结束": "2026-08-03 23:59:59",
    },
]


class AnalysisTest(unittest.TestCase):
    def test_export_and_analyze_sample_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook = export_workbook(SAMPLE_ROWS, tmp, filename="sample.xlsx")
            result = analyze_workbook(workbook)

            self.assertEqual(result.metrics["sheet_rows"]["魔方罗盘销售明细"], 2)
            self.assertEqual(result.metrics["compass"]["stock_risk_count"], 1)
            self.assertGreater(result.metrics["vipdirect"]["roi"], 8)
            self.assertEqual(result.metrics["bct_gift"]["used_people"], 25)
            self.assertEqual(result.metrics["bct_gift"]["usage_quantity"], 35)
            self.assertEqual(result.metrics["bct_gift"]["date_end"], "2026-08-03 23:59:59")
            self.assertEqual(result.metrics["bct_scene"]["touch_users"], 260)
            self.assertIn("爆款策略追踪数据洞察与策略建议", result.markdown)
            for heading in [
                "## 结论摘要",
                "## 动作信息与覆盖边界",
                "## 核心指标对比",
                "## 辅助指标校验",
                "## 终端商业指标",
                "## 投放明细校验",
                "## 权益使用效果",
                "## 效果标签与置信度解释",
                "## 后续建议",
                "## 数据缺口与补数提示",
            ]:
                self.assertIn(heading, result.markdown)
            for required_text in ["核心指标", "辅助", "终端", "效果标签", "置信度", "参考组", "活动周期"]:
                self.assertIn(required_text, result.markdown)
            self.assertTrue(result.recommendations)

            outputs = write_analysis(result, tmp)
            self.assertTrue(Path(outputs["analysis_json"]).is_file())
            self.assertTrue(Path(outputs["strategy_markdown"]).is_file())

    def test_cli_analyze_and_doc_summary_write_expected_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook = export_workbook(SAMPLE_ROWS, tmp, filename="sample.xlsx")

            stdout = StringIO()
            with redirect_stdout(stdout):
                exit_code = cli_main(["analyze", str(workbook), "--out-dir", tmp])
            self.assertEqual(exit_code, 0)
            self.assertIn("analysis_json", stdout.getvalue())
            self.assertTrue((Path(tmp) / "sample_strategy.md").is_file())

            stdout = StringIO()
            with redirect_stdout(stdout):
                exit_code = cli_main(["doc-summary", "--out-dir", tmp])
            self.assertEqual(exit_code, 0)
            self.assertIn("requirements-understanding.md", stdout.getvalue())
            self.assertTrue((Path(tmp) / "requirements-understanding.md").is_file())

    def test_missing_bct_fields_render_as_missing_not_zero(self) -> None:
        rows = [
            {"__sheet_name": "执行摘要", "报表来源": "test", "执行结果": "ok", "备注": "live-like missing bct fields"},
            {
                "__sheet_name": "魔方罗盘销售明细",
                "款号": "A100",
                "商品ID": "10001",
                "销售额": 50000,
                "客户数": 976,
                "商详UV": 13036,
                "曝光人数": 237172,
                "点击人数": 17195,
                "加购数": 546,
            },
            {
                "__sheet_name": "中台礼金",
                "活动ID": "G-live",
                "活动名称": "老友礼金",
                "运营方式": "老友礼金",
                "活动状态": "进行中",
                "已使用人数": 0,
                "销售额": 371701.68,
                "计划开始": "2026-07-29 10:00:00",
                "计划结束": "2026-08-06 19:59:59",
            },
            {
                "__sheet_name": "中台购物车跨品类券",
                "活动ID": "C-live",
                "活动名称": "购物车挽回",
                "运营方式": "购物车/跨品类券",
                "活动状态": "进行中",
                "已使用人数": 0,
                "销售额": 48958.56,
                "计划开始": "2026-07-29 10:00:00",
                "计划结束": "2026-08-06 19:59:59",
            },
        ]
        with tempfile.TemporaryDirectory() as tmp:
            workbook = export_workbook(rows, tmp, filename="missing_bct.xlsx")
            result = analyze_workbook(workbook)

            gift = result.metrics["bct_gift"]
            self.assertFalse(gift["touch_users_returned"])
            self.assertFalse(gift["usage_quantity_returned"])
            self.assertFalse(gift["orders_returned"])
            self.assertFalse(gift["cost_returned"])
            self.assertFalse(gift["roi_returned"])
            self.assertTrue(gift["used_people_returned"])
            self.assertEqual(gift["used_people"], 0)
            self.assertIn("礼金触达 未返回、使用人数 0、使用量 未返回", result.markdown)
            self.assertIn("活动成本 未返回，ROI 不可算", result.markdown)
            self.assertIn("报告按“未返回/不可算”展示", result.markdown)
            self.assertNotIn("活动成本 0.00，ROI 0.00", result.markdown)

    def test_bct_product_effect_attachment_expands_to_goods_rows(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "商品效果"
        ws.append(["导出说明", "", "", ""])
        ws.append(["商品ID", "货号", "商品名称", "销售额（元）", "ROI"])
        ws.append(["10001", "A100-RED", "红色T恤", 1234.5, 8.8])
        buffer = BytesIO()
        wb.save(buffer)
        encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
        rows = [
            {
                "__sheet_name": "中台礼金",
                "报表来源": "中台营销：礼金",
                "执行结果": "已下载商品效果文件",
                "__attachment_base64": encoded,
                "__attachment_file_name": "商品效果.xlsx",
                "__attachment_context": json.dumps({
                    "sheet": "中台礼金",
                    "sourceLabel": "中台营销：礼金",
                    "groupLabel": "用户生命周期运营",
                    "brandSn": "BALA",
                    "brandName": "巴拉巴拉",
                    "configKey": "toolTaskMidSummary",
                    "taskId": "D1",
                    "exportLabel": "商品汇总效果数据",
                }, ensure_ascii=False),
            }
        ]
        with tempfile.TemporaryDirectory() as tmp:
            expanded = _expand_download_attachments(rows, Path(tmp))
            self.assertEqual(len(expanded), 2)
            parsed = expanded[1]
            self.assertEqual(parsed["__sheet_name"], "中台礼金")
            self.assertEqual(parsed["商品ID"], "10001")
            self.assertEqual(parsed["货号"], "A100-RED")
            self.assertEqual(parsed["销售额"], 1234.5)
            self.assertEqual(parsed["下载中心配置"], "toolTaskMidSummary")
            self.assertTrue((Path(tmp) / "商品效果.xlsx").is_file())

            workbook = export_workbook(expanded, tmp, filename="expanded.xlsx")
            headers = [cell.value for cell in load_workbook(workbook)["中台礼金"][1]]
            self.assertNotIn("__attachment_base64", headers)

    def test_compass_backend_export_attachment_expands_and_filters_bala_rows(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "商品明细"
        ws.append(["导出说明", "", "", "", "", "", ""])
        ws.append(["品牌店铺", "MID", "货号", "商品名称", "销售额", "商详UV", "曝光量", "点击率", "在售库存"])
        ws.append(["巴拉巴拉Balabala", "20001", "BALA-RED", "童装T恤", 2888.0, 1200, 5000, "6.50%", 80])
        ws.append(["其它品牌", "30001", "OTHER-BLU", "其它商品", 1888.0, 900, 3000, "3.20%", 40])
        buffer = BytesIO()
        wb.save(buffer)
        encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
        rows = [
            {
                "__sheet_name": "魔方罗盘销售明细",
                "报表来源": "魔方罗盘销售明细",
                "执行结果": "已下载后台导出文件",
                "__attachment_base64": encoded,
                "__attachment_file_name": "compass.xlsx",
                "__attachment_context": json.dumps({
                    "sheet": "魔方罗盘销售明细",
                    "sourceLabel": "魔方罗盘销售明细",
                    "groupLabel": "商品明细",
                    "brandName": "巴拉巴拉",
                    "brandKeyword": "巴拉巴拉",
                    "balaOnly": True,
                    "configKey": "compass_product_detail_export",
                    "taskId": "C1",
                    "sourceApi": "/pc/mms/v1/item/export -> /pc/mms/v1/item/export/task/list",
                }, ensure_ascii=False),
            }
        ]
        with tempfile.TemporaryDirectory() as tmp:
            expanded = _expand_download_attachments(rows, Path(tmp))
            self.assertEqual(len(expanded), 2)
            parsed = expanded[1]
            self.assertEqual(parsed["__sheet_name"], "魔方罗盘销售明细")
            self.assertEqual(parsed["商品ID"], "20001")
            self.assertEqual(parsed["MID"], "20001")
            self.assertEqual(parsed["货号"], "BALA-RED")
            self.assertEqual(parsed["品牌名称"], "巴拉巴拉")
            self.assertEqual(parsed["品牌店铺"], "巴拉巴拉Balabala")
            self.assertEqual(parsed["曝光次数"], 5000)
            self.assertEqual(parsed["导出任务ID"], "C1")
            self.assertIn("/pc/mms/v1/item/export", parsed["数据来源接口"])

            workbook = export_workbook(expanded, tmp, filename="compass-expanded.xlsx")
            sheet = load_workbook(workbook)["魔方罗盘销售明细"]
            headers = [cell.value for cell in sheet[1]]
            self.assertIn("导出任务ID", headers)
            self.assertIn("源文件工作表", headers)
            mids = [row[headers.index("商品ID")].value for row in sheet.iter_rows(min_row=2)]
            self.assertIn("20001", mids)
            self.assertNotIn("30001", mids)

    def test_vipdirect_workbook_headers_include_goods_ad_detail_fields(self) -> None:
        rows = [
            {"__sheet_name": "执行摘要", "报表来源": "test", "执行结果": "ok", "备注": "vipdirect detail"},
            {
                "__sheet_name": "唯直达投放效果",
                "数据分组": "广告商品明细",
                "广告名称": "搜索广告A",
                "广告状态": "投放中",
                "商品ID": "20001",
                "MID": "20001",
                "货号": "BALA-RED",
                "商品名称": "童装T恤",
                "投放资源": "唯直达",
                "计费类型": "CPC",
                "花费": 100,
            },
        ]
        with tempfile.TemporaryDirectory() as tmp:
            workbook = export_workbook(rows, tmp, filename="vipdirect-detail.xlsx")
            headers = [cell.value for cell in load_workbook(workbook)["唯直达投放效果"][1]]
            for header in ["广告状态", "商品ID", "MID", "货号", "商品名称", "图片链接", "投放资源", "计费类型"]:
                self.assertIn(header, headers)
            self.assertLess(headers.index("商品ID"), headers.index("花费"))

    def test_compass_reported_percent_is_not_multiplied_twice(self) -> None:
        rows = [
            {"__sheet_name": "执行摘要", "报表来源": "test", "执行结果": "ok", "备注": "reported percent"},
            {
                "__sheet_name": "魔方罗盘销售明细",
                "款号": "A100",
                "商品ID": "10001",
                "销售额": 10000,
                "客户数": 10,
                "商详UV": 100,
                "点击人数": 25,
                "商品点击率": 7.25,
            },
        ]
        with tempfile.TemporaryDirectory() as tmp:
            workbook = export_workbook(rows, tmp, filename="percent.xlsx")
            result = analyze_workbook(workbook)

            self.assertAlmostEqual(result.metrics["compass"]["click_rate"], 0.0725)
            self.assertIn("商品点击率 7.25%", result.markdown)
            self.assertNotIn("商品点击率 725.00%", result.markdown)

    def test_collect_help_exposes_url_prefix_selector(self) -> None:
        stdout = StringIO()
        with self.assertRaises(SystemExit) as captured:
            with redirect_stdout(stdout):
                cli_main(["collect", "--help"])
        self.assertEqual(captured.exception.code, 0)
        self.assertIn("--url-prefix", stdout.getvalue())
        self.assertIn("--bct-gift-activity-start", stdout.getvalue())
        self.assertIn("--bct-scene-activity-end", stdout.getvalue())
        self.assertIn("--compass-table-api", stdout.getvalue())
        self.assertIn("--vipdirect-summary-api", stdout.getvalue())

    def test_collect_params_keep_independent_report_ranges(self) -> None:
        args = SimpleNamespace(
            page_size=100,
            bct_page_size=20,
            max_pages=1,
            brand_keyword="巴拉巴拉",
            compass_start_date="2026-08-02",
            compass_end_date="2026-08-02",
            vipdirect_start_date="2026-08-01",
            vipdirect_end_date="2026-08-02",
            tmax_start_date="2026-07-30",
            tmax_end_date="2026-08-02",
            bct_activity_start="",
            bct_activity_end="",
            bct_gift_activity_start="2026-07-29",
            bct_gift_activity_end="2026-08-06",
            bct_scene_activity_start="2026-08-01",
            bct_scene_activity_end="2026-08-03",
            compass_export_wait_seconds=0,
            vipdirect_detail_source="ad_list",
        )
        params = build_params(args)

        self.assertEqual(params["compass_start_date"], "2026-08-02")
        self.assertEqual(params["compass_export_wait_seconds"], 0)
        self.assertEqual(params["vipdirect_end_date"], "2026-08-02")
        self.assertEqual(params["vipdirect_detail_source"], "ad_list")
        self.assertEqual(params["tmax_start_date"], "2026-07-30")
        self.assertEqual(params["bct_gift_activity_start"], "2026-07-29")
        self.assertEqual(params["bct_scene_activity_end"], "2026-08-03")


if __name__ == "__main__":
    unittest.main()
