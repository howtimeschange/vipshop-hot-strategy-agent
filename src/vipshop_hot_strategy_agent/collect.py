from __future__ import annotations

import asyncio
import base64
import csv
import importlib.util
import json
import sys
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from dataclasses import replace
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Callable, Mapping

from openpyxl import load_workbook

from .excel_export import export_workbook
from .paths import paths


DEFAULT_SCOPE = ["compass_sales_detail", "vipdirect_ads", "tmax_goods", "bct_gift", "bct_scene"]


def runtime_candidates() -> list[Path]:
    return [
        Path("/Users/xingyicheng/.codex/plugins/cache/personal/crawshrimp-skill/0.1.0+codex.20260622090512/skills/crawshrimp-skill/scripts"),
        Path("/Users/xingyicheng/Documents/crawshrimp/skills/crawshrimp-skill/scripts"),
        paths().root / "vendor" / "crawshrimp-skill" / "scripts",
    ]


def locate_runtime_scripts(runtime_root: str | Path | None = None) -> Path:
    candidates = [Path(runtime_root).expanduser()] if runtime_root else runtime_candidates()
    for candidate in candidates:
        if (candidate / "browser_executor.py").is_file() and (candidate / "phase_runner.py").is_file():
            return candidate
    searched = "\n".join(str(item) for item in candidates)
    raise FileNotFoundError(f"Cannot locate crawshrimp-skill scripts. Searched:\n{searched}")


def _load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise ImportError(f"Cannot load {name} from {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def load_runtime(runtime_root: str | Path | None = None):
    scripts = locate_runtime_scripts(runtime_root)
    if str(scripts) not in sys.path:
        sys.path.insert(0, str(scripts))
    browser_executor = _load_module("browser_executor", scripts / "browser_executor.py")
    phase_runner = _load_module("phase_runner", scripts / "phase_runner.py")
    return browser_executor, phase_runner


def build_params(args: Any) -> dict[str, Any]:
    params: dict[str, Any] = {
        "report_scope": list(DEFAULT_SCOPE),
        "page_size": args.page_size,
        "bct_page_size": args.bct_page_size,
        "max_pages": args.max_pages,
    }
    for key in [
        "brand_keyword",
        "compass_start_date",
        "compass_end_date",
        "compass_brand_store_sn",
        "compass_table_api",
        "compass_table_fallback",
        "compass_export_required",
        "compass_force_new_export",
        "compass_export_wait_seconds",
        "compass_export_poll_seconds",
        "compass_export_submit_url",
        "compass_export_task_list_url",
        "vipdirect_start_date",
        "vipdirect_end_date",
        "vipdirect_summary_api",
        "vipdirect_detail_source",
        "vipdirect_mid",
        "vipdirect_no_product_insight_fallback",
        "tmax_start_date",
        "tmax_end_date",
        "bct_activity_start",
        "bct_activity_end",
        "bct_gift_activity_start",
        "bct_gift_activity_end",
        "bct_scene_activity_start",
        "bct_scene_activity_end",
        "bct_product_export_type",
        "bct_download_wait_seconds",
        "bct_download_poll_seconds",
    ]:
        value = getattr(args, key, None)
        if value is not None and value != "":
            params[key] = value
    return params


DOWNLOAD_ROW_ALIASES: dict[str, tuple[str, ...]] = {
    "品牌编码": ("品牌编码", "品牌SN", "品牌编号", "brandStoreSn", "brand_store_sn", "brandSn"),
    "品牌名称": ("品牌名称", "品牌", "品牌店铺", "brandStoreName", "brand_store_name", "brandName"),
    "统计日期": ("统计日期", "日期", "数据日期", "dt", "date", "statDate", "stat_date"),
    "MID": ("MID", "mid", "商品MID", "商品 Mid", "商品id", "商品ID", "merchandiseNo", "merchandise_no"),
    "商品ID": ("商品ID", "商品id", "商品Id", "商品编号", "商品MID", "MID", "mid", "spuId", "SPU ID", "spu_id", "merchandiseNo", "merchandise_no"),
    "款号": ("款号", "大货款号", "styleCode", "style_code", "styleNo", "osn"),
    "货号": ("货号", "商品货号", "商品编码", "goodsNo", "goods_no", "msn", "sn", "outerId"),
    "商品名称": ("商品名称", "商品名", "货品名称", "goodsName", "goods_name", "productName", "product_name"),
    "P_SPU_ID": ("P_SPU_ID", "P_SPU ID", "PSPU_ID", "prodSpuId", "prod_spu_id"),
    "商品图片": ("商品图片", "图片链接", "图片", "imgUrl", "imageUrl", "picUrl"),
    "一级品类": ("一级品类", "一级类目", "firstCateName", "first_cate_name"),
    "二级品类": ("二级品类", "二级类目", "secCateName", "secondCateName", "sec_cate_name"),
    "三级品类": ("三级品类", "三级类目", "三级分类名称", "thirdCateName", "third_cate_name"),
    "售卖渠道": ("售卖渠道", "渠道", "channelName", "channelType"),
    "售龄": ("售龄", "productSellAge"),
    "货品类型": ("货品类型", "商品标签", "标签", "merTypeList"),
    "适用季节": ("适用季节", "suitSeason"),
    "市场价": ("市场价", "最低市场价", "市场价（元）", "市场价(元)", "minMarketPrice"),
    "唯品价": ("唯品价", "最低唯品价", "唯品价（元）", "唯品价(元)", "minVipshopPrice"),
    "到手价": ("到手价", "最低到手价", "到手价（元）", "到手价(元)", "minPayPrice"),
    "折扣": ("折扣", "折扣区间", "discountRangeName", "discountRange"),
    "首次上架时间": ("首次上架时间", "首上时间", "firstShelfTime"),
    "主款式": ("主款式", "mainStyleName"),
    "是否上架": ("是否上架", "商品状态", "状态", "status", "statusName"),
    "曝光UV": ("曝光UV", "曝光uv", "曝光用户数", "曝光人数", "impressionUv", "exposureUserNum", "exposeUserNum"),
    "商详UV": ("商详UV", "商品详情页UV", "详情UV", "uv", "detailUv", "goodsUv"),
    "曝光-点击率": ("曝光-点击率", "曝光-点击率 (商详UV/曝光UV)", "曝光点击率", "商品点击率", "点击率", "impressionClickRate"),
    "CTR": ("CTR", "CTR(点击PV/曝光PV)", "ctr"),
    "商详UV价值": ("商详UV价值", "商详UV价值(销售额/商详UV)", "valueOfUv"),
    "曝光UV价值": ("曝光UV价值", "曝光UV价值(销售额/曝光UV)", "valueOfImpressionUv"),
    "收藏人数": ("收藏人数", "收藏数", "collectNum", "favNum", "favoriteNum", "merchandiseLikeUserNum"),
    "收藏率": ("收藏率", "收藏率(收藏人数/商详UV)", "merchandiseLikeRate"),
    "加购人数": ("加购人数", "加购UV(加购用户数)", "加购数", "cartNum", "addCartNum", "addToCartNum", "addUserNum"),
    "加购率": ("加购率", "加购率(加购人数/商详UV)", "addRate"),
    "兴趣人数": ("兴趣人数", "interestedUserNum"),
    "兴趣人次": ("兴趣人次", "interestedUserMerNum"),
    "销售额": ("销售额", "销售额（元）", "销售额(元)", "用券销售额", "使用销售额", "购买金额", "GMV", "buyAmt", "buyTotal", "goodsActureAmt", "actureAmt", "saleAmt", "salesAmt", "payAmt"),
    "销售量": ("销售量", "销售数量", "销量", "销售件数", "goodsActureNum", "actureNum", "saleNum", "salesNum"),
    "销售数量": ("销售数量", "销量", "销售件数", "goodsActureNum", "actureNum", "saleNum", "salesNum"),
    "订单额": ("订单额", "订单金额", "orderAmt", "orderAmount"),
    "客户数": ("客户数", "购买用户数", "用券客户数", "useUserNum", "customNum", "userNum", "customerNum"),
    "品牌新客数": ("品牌新客数", "bctNewUser"),
    "购买转化率": ("购买转化率", "购买转化率(客户数/商详UV)", "pageMerDetailConvRate"),
    "优惠券金额": ("优惠券金额", "couponAmt"),
    "销售额(不含拒退)": ("销售额(不含拒退)", "销售额（不含拒退）", "resultGoodsActureAmt"),
    "销售量(不含拒退)": ("销售量(不含拒退)", "销售数量(不含拒退)", "销售量（不含拒退）", "销售数量（不含拒退）", "resultGoodsActureNum"),
    "件单价": ("件单价", "件单价(销售额/销售量)", "itemUnitPrice"),
    "ARPU": ("ARPU", "ARPU(销售额/客户数)", "arpu", "userUnitPrice"),
    "超V销售额": ("超V销售额", "svipGoodsActureAmt"),
    "超V销售量": ("超V销售量", "svipGoodsActureNum"),
    "超V客户数": ("超V客户数", "svipUserNum"),
    "品牌会员销售额": ("品牌会员销售额", "brandGoodsActureAmt"),
    "品牌会员销售量": ("品牌会员销售量", "brandGoodsActureNum"),
    "品牌会员客户数": ("品牌会员客户数", "brandUserNum"),
    "退货件数": ("退货件数", "returnItemNum"),
    "拒收件数": ("拒收件数", "rejectItemNum"),
    "拒退件数": ("拒退件数", "rejectAndReturnItemNum"),
    "换货件数": ("换货件数", "exchangeItemNum"),
    "退货金额": ("退货金额", "returnGoodsAmt"),
    "拒收金额": ("拒收金额", "rejectGoodsAmt"),
    "拒退金额": ("拒退金额", "rejectAndReturnGoodsAmt"),
    "换货金额": ("换货金额", "exchangeGoodsAmt"),
    "退货率": ("退货率", "退货率(退货件数/销售量)", "returnItemRate"),
    "拒收率": ("拒收率", "拒收率(拒收件数/销售量)", "rejectItemRate"),
    "拒退率": ("拒退率", "rejectAndReturnItemRate"),
    "换货率": ("换货率", "换货率(换货件数/销售量)", "exchangeItemRate"),
    "在售库存": ("在售库存", "库存", "onSellLeavingNum"),
    "可售库存": ("可售库存", "leavingNum"),
    "在售唯品价货值": ("在售唯品价货值", "onSellValueOfVipshop"),
    "在售到手价货值": ("在售到手价货值", "onSellValueOfPay"),
    "售卖比": ("售卖比", "货量售卖比", "salesRateOfNum"),
    "用户策略ROI": ("用户策略ROI", "userStrategyRoi"),
    "用户策略销售额": ("用户策略销售额", "userStrategyGoodsActureAmt"),
    "用户策略成本": ("用户策略成本", "userStrategyCostAmt"),
    "会员权益ROI": ("会员权益ROI", "brandUserStrategyRoi"),
    "会员权益销售额": ("会员权益销售额", "brandUserStrategyGoodsActureAmt"),
    "会员权益成本": ("会员权益成本", "brandUserStrategyCostAmt"),
    "唯享客ROI": ("唯享客ROI", "wxkStrategyRoi"),
    "唯享客销售额": ("唯享客销售额", "wxkStrategyGoodsActureAmt"),
    "唯享客成本": ("唯享客成本", "wxkStrategyCostAmt"),
    "曝光人数": ("曝光人数", "曝光用户数", "exposureUserNum", "exposeUserNum"),
    "曝光次数": ("曝光次数", "曝光量", "曝光PV", "impressionNum", "impressionCount"),
    "访问人数": ("访问人数", "访问用户数", "visitorNum", "returnUserNum"),
    "点击人数": ("点击人数", "点击用户数", "clickUserNum", "clickNum"),
    "商品点击率": ("商品点击率", "点击率", "impressionClickRate", "ctr"),
    "收藏数": ("收藏数", "收藏人数", "collectNum", "favNum", "favoriteNum", "merchandiseLikeUserNum"),
    "加购数": ("加购数", "加购人数", "cartNum", "addCartNum", "addToCartNum", "addUserNum"),
    "加购件数": ("加购件数", "加购商品件数", "addItemNum"),
    "下单客户数": ("下单客户数", "下单用户数", "orderUserNum", "orderCustomerNum"),
    "库存天数": ("库存天数", "可售库存天数", "canSellStockDay"),
    "商品状态": ("商品状态", "状态", "status", "statusName"),
    "商品标签": ("商品标签", "标签", "merTypeList"),
    "图片链接": ("图片链接", "商品图片", "图片", "imgUrl", "imageUrl", "picUrl"),
    "活动ID": ("活动ID", "活动id", "活动编号", "任务ID", "任务id", "策略ID", "id"),
    "活动名称": ("活动名称", "任务名称", "策略名称", "名称", "taskName"),
    "人群名称": ("人群名称", "uspGroupName"),
    "运营方式": ("运营方式", "类型", "任务类型", "taskTypeName"),
    "营销渠道": ("营销渠道", "taskChannels"),
    "活动状态": ("活动状态", "状态", "statusName"),
    "计划开始": ("计划开始", "活动开始时间", "开始时间", "startTime"),
    "计划结束": ("计划结束", "活动结束时间", "结束时间", "endTime"),
    "触达用户数": ("触达用户数", "触达人数", "touchUserNum", "marketUserNum"),
    "已使用人数": ("已使用人数", "使用人数", "用券人数", "useUser", "usedUserNum"),
    "使用量": ("使用量", "使用次数", "用券量", "couponUseNum", "useNum"),
    "用券量": ("用券量", "使用量", "使用次数", "couponUseNum", "useNum"),
    "订单数": ("订单数", "订单量", "购买订单数", "orderNum", "saleNum"),
    "客单价": ("客单价", "客单价（元）", "客单价(元)", "arpu", "unitPrice"),
    "活动成本": ("活动成本", "活动成本（元）", "活动成本(元)", "成本", "cost", "activityCost"),
    "ROI": ("ROI", "roi"),
}

HEADER_KEYWORDS = ("商品", "货号", "款号", "活动", "销售", "订单", "客户", "ROI", "成本", "曝光", "访问", "MID", "品牌")


def _visible_row(row: Mapping[str, Any]) -> dict[str, Any]:
    return {
        key: value
        for key, value in row.items()
        if key == "__sheet_name" or not str(key).startswith("__")
    }


def _is_empty(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _first_present(row: Mapping[str, Any], keys: tuple[str, ...]) -> Any:
    lowered = {str(key).strip().lower(): value for key, value in row.items()}
    for key in keys:
        for candidate in (key, key.lower()):
            if candidate in row and not _is_empty(row[candidate]):
                return row[candidate]
            lowered_value = lowered.get(str(candidate).strip().lower())
            if not _is_empty(lowered_value):
                return lowered_value
    return ""


def _dedupe_headers(headers: list[str]) -> list[str]:
    counts: dict[str, int] = {}
    output: list[str] = []
    for index, header in enumerate(headers, start=1):
        name = str(header or "").strip() or f"列{index}"
        counts[name] = counts.get(name, 0) + 1
        output.append(name if counts[name] == 1 else f"{name}_{counts[name]}")
    return output


def _header_score(values: list[Any]) -> int:
    labels = [str(value or "").strip() for value in values]
    non_empty = [label for label in labels if label]
    if len(non_empty) < 2:
        return 0
    return len(non_empty) + sum(3 for label in non_empty if any(keyword in label for keyword in HEADER_KEYWORDS))


def _worksheet_rows(ws: Any) -> list[dict[str, Any]]:
    if hasattr(ws, "reset_dimensions"):
        ws.reset_dimensions()
    raw_rows = [list(row) for row in ws.iter_rows(values_only=True)]
    if not raw_rows:
        return []
    best_index = -1
    best_score = 0
    for index, row in enumerate(raw_rows[:20]):
        score = _header_score(row)
        if score > best_score:
            best_index = index
            best_score = score
    if best_index < 0:
        return []
    headers = _dedupe_headers([str(value or "").strip() for value in raw_rows[best_index]])
    rows: list[dict[str, Any]] = []
    for raw in raw_rows[best_index + 1:]:
        values = list(raw) + [""] * (len(headers) - len(raw))
        if all(_is_empty(value) for value in values):
            continue
        rows.append({headers[index]: values[index] for index in range(len(headers))})
    return rows


def _csv_rows(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig")
    rows = list(csv.reader(StringIO(text)))
    if not rows:
        return []
    best_index = max(range(min(20, len(rows))), key=lambda index: _header_score(rows[index]))
    headers = _dedupe_headers(rows[best_index])
    output: list[dict[str, Any]] = []
    for raw in rows[best_index + 1:]:
        values = list(raw) + [""] * (len(headers) - len(raw))
        if all(_is_empty(value) for value in values):
            continue
        output.append({headers[index]: values[index] for index in range(len(headers))})
    return output


def _safe_file_name(name: str, fallback: str) -> str:
    cleaned = "".join(char if char.isalnum() or char in "._-（）()[] " else "_" for char in str(name or "").strip())
    return cleaned[:160] or fallback


def _filename_from_content_disposition(value: str) -> str:
    if not value:
        return ""
    star = "filename*=UTF-8''"
    for part in value.split(";"):
        text = part.strip()
        if text.lower().startswith(star.lower()):
            return urllib.parse.unquote(text[len(star):].strip().strip('"'))
        if text.lower().startswith("filename="):
            return text.split("=", 1)[1].strip().strip('"')
    return ""


def _suffix_for_attachment(data: bytes, content_type: str) -> str:
    lowered = content_type.lower()
    if "csv" in lowered:
        return ".csv"
    if "spreadsheetml" in lowered or "excel" in lowered:
        return ".xlsx"
    if data.startswith(b"PK\x03\x04"):
        return ".xlsx"
    return ""


def _decode_attachment_context(row: Mapping[str, Any]) -> dict[str, Any]:
    raw = row.get("__attachment_context") or "{}"
    if isinstance(raw, Mapping):
        return dict(raw)
    try:
        parsed = json.loads(str(raw))
        return parsed if isinstance(parsed, dict) else {}
    except json.JSONDecodeError:
        return {}


def _request_attachment(
    url: str,
    *,
    method: str = "GET",
    data: bytes | None = None,
    content_type: str = "",
) -> tuple[bytes, str, str]:
    headers = {"User-Agent": "vipshop-hot-strategy-agent/0.1"}
    if content_type:
        headers["Content-Type"] = content_type
    request = urllib.request.Request(url, data=data, headers=headers, method=method)
    with urllib.request.urlopen(request, timeout=120) as response:
        body = response.read()
        response_type = response.headers.get("content-type", "")
        file_name = _filename_from_content_disposition(response.headers.get("content-disposition", ""))
        return body, file_name, response_type


def _looks_like_download(data: bytes, content_type: str) -> bool:
    lowered = content_type.lower()
    if data.startswith(b"PK\x03\x04") or data.startswith(b"\xef\xbb\xbf"):
        return True
    if any(marker in lowered for marker in ["spreadsheet", "excel", "csv", "octet-stream", "force-download"]):
        return True
    stripped = data.lstrip()[:20].lower()
    return not any(marker in lowered for marker in ["application/json", "text/html"]) and not stripped.startswith((b"{", b"[", b"<"))


def _parse_download_payload(payload_text: str) -> Any:
    if not payload_text:
        return {}
    try:
        return json.loads(payload_text)
    except json.JSONDecodeError:
        return payload_text


def _payload_query(payload: Any) -> str:
    if isinstance(payload, Mapping):
        return urllib.parse.urlencode({str(key): value for key, value in payload.items()}, doseq=True)
    return str(payload or "")


def _download_attachment_bytes(row: Mapping[str, Any]) -> tuple[bytes, str, str]:
    url = str(row.get("__attachment_download_url") or "").strip()
    if not url:
        return b"", "", ""
    method = str(row.get("__attachment_download_method") or "GET").strip().upper() or "GET"
    payload_text = str(row.get("__attachment_download_payload") or "").strip()
    if method != "POST":
        try:
            return _request_attachment(url)
        except urllib.error.URLError as error:
            raise RuntimeError(f"下载链接读取失败：{error}") from error

    payload = _parse_download_payload(payload_text)
    attempts: list[tuple[str, str, bytes | None, str]] = []
    attempts.append(("JSON", url, (payload_text or "{}").encode("utf-8"), "application/json;charset=UTF-8"))
    query_text = _payload_query(payload)
    if query_text:
        attempts.append(("FormData", url, query_text.encode("utf-8"), "application/x-www-form-urlencoded;charset=UTF-8"))
        if isinstance(payload, Mapping) and "reqData" not in payload:
            attempts.append((
                "reqData",
                url,
                urllib.parse.urlencode({"reqData": json.dumps(payload, ensure_ascii=False)}).encode("utf-8"),
                "application/x-www-form-urlencoded;charset=UTF-8",
            ))
        attempts.append(("query", f"{url}{'&' if '?' in url else '?'}{query_text}", None, ""))

    errors: list[str] = []
    fallback: tuple[bytes, str, str] | None = None
    for label, attempt_url, data, content_type in attempts:
        try:
            result = _request_attachment(attempt_url, method="POST" if data is not None else "GET", data=data, content_type=content_type)
            if _looks_like_download(result[0], result[2]):
                return result
            fallback = result
            errors.append(f"{label}: 响应不是文件（{result[2] or 'unknown content-type'}）")
        except urllib.error.URLError as error:
            errors.append(f"{label}: {error}")
    if fallback:
        return fallback
    raise RuntimeError(f"下载链接读取失败：{'；'.join(errors)}")


def _context_flag(context: Mapping[str, Any], key: str, default: bool = False) -> bool:
    value = context.get(key)
    if value in (None, ""):
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "y", "是"}


def _download_row_matches_context(row: Mapping[str, Any], context: Mapping[str, Any]) -> bool:
    sheet = str(context.get("sheet") or "").strip()
    if sheet != "魔方罗盘销售明细" or not _context_flag(context, "balaOnly", False):
        return True
    keyword = str(context.get("brandKeyword") or context.get("brandName") or "巴拉巴拉").strip()
    context_brand = str(context.get("brandName") or "").strip()
    brand_values = [
        str(value).strip()
        for key in ["品牌店铺", "品牌", "brandStoreName", "brandName", "brand_store_name"]
        if not _is_empty(value := row.get(key))
    ]
    brand_name = str(row.get("品牌名称") or "").strip()
    if brand_name and (not context_brand or brand_name != context_brand):
        brand_values.append(brand_name)
    haystack = " ".join(brand_values)
    if not haystack:
        return True
    return (keyword and keyword in haystack) or ("巴拉巴拉" in haystack) or ("Balabala".lower() in haystack.lower())


def _normalize_download_row(
    raw_row: Mapping[str, Any],
    *,
    context: Mapping[str, Any],
    file_name: str,
    saved_path: Path,
    source_sheet: str,
) -> dict[str, Any]:
    sheet = str(context.get("sheet") or "中台礼金")
    source_api = context.get("sourceApi") or "/bcc/vendorBrand/downCenter/add -> /bcc/vendorBrand/downCenter/page"
    result_text = "已解析后台导出" if sheet == "魔方罗盘销售明细" else "已解析商品效果"
    remark_text = "来源：魔方后台下载导出商品明细文件" if sheet == "魔方罗盘销售明细" else "来源：右上角「批量导出商品效果」下载文件"
    row: dict[str, Any] = {
        "__sheet_name": sheet,
        "报表来源": context.get("sourceLabel") or "",
        "数据分组": context.get("groupLabel") or "",
        "品牌编码": context.get("brandSn") or "",
        "品牌名称": context.get("brandName") or "",
        "供应商编码": context.get("vendorCode") or "",
        "商品效果导出类型": context.get("exportLabel") or "",
        "下载中心配置": context.get("configKey") or "",
        "下载中心任务ID": context.get("taskId") or "",
        "导出任务ID": context.get("taskId") or "",
        "下载中心状态": "生成成功",
        "下载文件名": file_name,
        "下载文件路径": str(saved_path),
        "源文件工作表": source_sheet,
        "数据来源接口": source_api,
        "执行结果": result_text,
        "备注": remark_text,
    }
    for key, value in raw_row.items():
        if not _is_empty(value):
            row[str(key).strip()] = value
    for target, aliases in DOWNLOAD_ROW_ALIASES.items():
        if _is_empty(row.get(target)):
            value = _first_present(raw_row, aliases)
            if not _is_empty(value):
                row[target] = value
    return row


def _parse_download_attachment(row: Mapping[str, Any], out_dir: Path) -> list[dict[str, Any]]:
    encoded = str(row.get("__attachment_base64") or "")
    downloaded_name = ""
    downloaded_type = ""
    if encoded:
        data = base64.b64decode(encoded)
    else:
        data, downloaded_name, downloaded_type = _download_attachment_bytes(row)
    if not data:
        return []
    provided_name = str(row.get("__attachment_file_name") or "").strip()
    file_name = _safe_file_name(provided_name or downloaded_name or "bct-product-effect.xlsx", "bct-product-effect.xlsx")
    effective_content_type = str(row.get("__attachment_content_type") or downloaded_type or "")
    if not Path(file_name).suffix:
        file_name = f"{file_name}{_suffix_for_attachment(data, effective_content_type) or '.xlsx'}"
    if downloaded_type and _is_empty(row.get("__attachment_content_type")):
        row = dict(row)
        row["__attachment_content_type"] = downloaded_type
    saved_path = out_dir / file_name
    counter = 1
    while saved_path.exists():
        saved_path = out_dir / f"{Path(file_name).stem}-{counter}{Path(file_name).suffix}"
        counter += 1
    saved_path.write_bytes(data)
    context = _decode_attachment_context(row)
    if file_name.lower().endswith(".csv") or "csv" in str(row.get("__attachment_content_type") or "").lower():
        output: list[dict[str, Any]] = []
        for item in _csv_rows(data):
            normalized = _normalize_download_row(item, context=context, file_name=file_name, saved_path=saved_path, source_sheet="CSV")
            if _download_row_matches_context(normalized, context):
                output.append(normalized)
        return output
    workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    output: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        for item in _worksheet_rows(worksheet):
            normalized = _normalize_download_row(
                item,
                context=context,
                file_name=file_name,
                saved_path=saved_path,
                source_sheet=worksheet.title,
            )
            if _download_row_matches_context(normalized, context):
                output.append(normalized)
    return output


def _expand_download_attachments(rows: list[Mapping[str, Any]], out_dir: Path) -> list[dict[str, Any]]:
    expanded: list[dict[str, Any]] = []
    for row in rows:
        visible = _visible_row(row)
        if not row.get("__attachment_base64") and not row.get("__attachment_download_url"):
            expanded.append(dict(row))
            continue
        try:
            parsed_rows = _parse_download_attachment(row, out_dir)
        except Exception as error:  # noqa: BLE001 - keep collection truthful instead of dropping the task row.
            visible["执行结果"] = "商品效果文件解析失败"
            visible["备注"] = f"{visible.get('备注') or ''}；解析失败：{error}".strip("；")
            expanded.append(visible)
            continue
        visible["执行结果"] = "已解析商品效果文件"
        visible["备注"] = f"{visible.get('备注') or ''}；已解析 {len(parsed_rows)} 行商品效果".strip("；")
        expanded.append(visible)
        expanded.extend(parsed_rows)
    return expanded


async def collect_reports_async(
    *,
    cdp_url: str = "http://127.0.0.1:9222",
    tab_id: str = "",
    url_prefix: str = "",
    runtime_root: str | Path | None = None,
    params: dict[str, Any] | None = None,
    out_dir: str | Path = "~/Downloads",
    progress: Callable[[dict[str, Any]], None] | None = None,
) -> dict[str, Any]:
    browser_executor, phase_runner = load_runtime(runtime_root)

    class LongEvalBackend(browser_executor.ChromeCDPBackend):
        async def execute_async(self, action):
            if action.kind == "eval":
                action = replace(action, timeout_ms=300000)
            return await super().execute_async(action)

    script_path = paths().assets / "hot-strategy-tracking-report.js"
    script = script_path.read_text(encoding="utf-8")
    backend = LongEvalBackend(cdp_url=cdp_url, tab_id=tab_id, url_prefix=url_prefix)
    runner = phase_runner.WebPhaseRunner(backend=backend, artifact_dir=paths().root / "artifacts")

    async def hook(event: dict[str, Any]) -> None:
        if progress:
            progress(event)

    result = await runner.run_script(script, params=params or {}, control_hook=hook)
    rows = result.data
    if not rows:
        raise RuntimeError("collection returned no rows")
    out_path = Path(out_dir).expanduser()
    out_path.mkdir(parents=True, exist_ok=True)
    rows = _expand_download_attachments(rows, out_path)
    workbook = export_workbook(rows, out_path)
    counts = Counter(str(row.get("__sheet_name") or "Sheet1") for row in rows)
    return {
        "workbook": str(workbook),
        "rows": len(rows),
        "sheet_counts": dict(counts),
        "shared": result.shared,
    }


def collect_reports(**kwargs: Any) -> dict[str, Any]:
    return asyncio.run(collect_reports_async(**kwargs))


def print_json(payload: dict[str, Any]) -> None:
    print(json.dumps(payload, ensure_ascii=False, indent=2))
