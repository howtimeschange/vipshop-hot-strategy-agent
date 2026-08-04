from __future__ import annotations

from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SUMMARY_SHEET = "执行摘要"
SHEET_ORDER = [
    SUMMARY_SHEET,
    "魔方罗盘销售明细",
    "中台礼金",
    "中台购物车跨品类券",
    "唯直达投放效果",
    "T-max效果",
]

SHEET_COLUMNS: dict[str, list[str]] = {
    SUMMARY_SHEET: ["报表来源", "执行结果", "备注"],
    "魔方罗盘销售明细": [
        "报表来源", "数据分组", "品牌编码", "品牌名称", "统计日期", "商品ID", "商品名称", "货号", "款号",
        "P_SPU_ID", "商品图片", "三级品类", "售龄", "货品类型", "适用季节", "唯品价", "到手价", "市场价",
        "折扣", "首次上架时间", "主款式", "是否上架", "曝光UV", "商详UV", "曝光-点击率", "CTR",
        "商详UV价值", "曝光UV价值", "收藏人数", "收藏率", "加购人数", "加购率", "兴趣人数", "兴趣人次",
        "销售额", "销售量", "客户数", "品牌新客数", "购买转化率", "优惠券金额",
        "销售额(不含拒退)", "销售量(不含拒退)", "件单价", "ARPU", "超V销售额", "超V销售量", "超V客户数",
        "品牌会员销售额", "品牌会员销售量", "品牌会员客户数", "退货件数", "拒收件数", "拒退件数", "换货件数",
        "退货金额", "拒收金额", "拒退金额", "换货金额", "退货率", "拒收率", "拒退率", "换货率",
        "在售库存", "可售库存", "在售唯品价货值", "在售到手价货值", "售卖比", "用户策略ROI",
        "用户策略销售额", "用户策略成本", "会员权益ROI", "会员权益销售额", "会员权益成本",
        "唯享客ROI", "唯享客销售额", "唯享客成本", "导出任务ID",
        "下载文件名", "下载文件类型", "下载文件字节数", "下载文件路径", "源文件工作表", "数据来源接口", "执行结果", "备注",
    ],
    "中台礼金": [
        "报表来源", "数据分组", "品牌编码", "品牌名称", "供应商编码", "商品效果导出类型", "下载中心配置",
        "下载中心任务ID", "下载中心状态", "下载文件名", "下载文件类型", "下载文件字节数", "下载文件路径",
        "筛选策略类型", "筛选条件", "源文件工作表", "选中活动ID", "商品ID", "款号", "货号", "商品名称",
        "一级品类", "二级品类", "三级品类", "活动ID", "联合提报ID", "活动名称",
        "人群名称", "运营方式", "营销渠道", "活动状态", "计划开始", "计划结束", "创建时间", "下线时间",
        "下线人", "券ID", "推广ID", "触达用户数", "曝光人数", "访问人数", "客户数", "已使用人数",
        "使用量", "销售额", "订单数", "客单价", "活动成本", "ROI", "数据来源接口", "执行结果", "备注",
    ],
    "中台购物车跨品类券": [
        "报表来源", "数据分组", "品牌编码", "品牌名称", "供应商编码", "商品效果导出类型", "下载中心配置",
        "下载中心任务ID", "下载中心状态", "下载文件名", "下载文件类型", "下载文件字节数", "下载文件路径",
        "筛选策略类型", "筛选条件", "源文件工作表", "选中活动ID", "商品ID", "款号", "货号", "商品名称",
        "一级品类", "二级品类", "三级品类", "活动ID", "联合提报ID", "活动名称",
        "人群名称", "运营方式", "营销渠道", "活动状态", "计划开始", "计划结束", "创建时间", "下线时间",
        "下线人", "券ID", "推广ID", "触达用户数", "曝光人数", "访问人数", "客户数", "已使用人数",
        "用券量", "销售额", "订单数", "客单价", "活动成本", "ROI", "数据来源接口", "执行结果", "备注",
    ],
    "唯直达投放效果": [
        "报表来源", "数据分组", "品牌名称", "统计日期", "活动ID", "活动名称", "推广ID", "广告名称", "广告状态",
        "商品ID", "MID", "货号", "商品名称", "图片链接", "关键词", "投放资源", "营销渠道", "计费类型",
        "花费", "曝光量", "点击量", "点击率", "千次曝光均价", "点击均价", "APP端UV", "小程序UV",
        "商品详情页UV", "品牌UV", "24小时收藏数(商家)", "24小时加购数(商家)", "24小时加购成本(商家)",
        "24小时下单客户数(商家)", "24小时下单量(商家)", "24小时订单额(商家)", "24小时下单ROI(商家)",
        "24小时成交客户数(商家)", "24小时销售额(商家)", "24小时成交ROI(商家)", "24小时收藏数(商品)",
        "24小时加购数(商品)", "24小时加购成本(商品)", "24小时下单客户数(商品)", "24小时下单量(商品)",
        "24小时订单额(商品)", "24小时下单ROI(商品)", "24小时成交客户数(商品)", "24小时销售额(商品)",
        "24小时成交ROI(商品)", "数据来源接口", "执行结果", "备注",
    ],
    "T-max效果": [
        "报表来源", "数据分组", "品牌名称", "活动ID", "活动名称", "推广ID", "商品ID", "货号", "商品名称",
        "运营方式", "营销渠道", "出价", "行动出价", "计费类型", "曝光人数", "点击人数", "收藏数", "加购数",
        "加购成本", "销售额", "ROI", "活动成本", "数据来源接口", "执行结果", "备注",
    ],
}


def timestamp() -> str:
    return datetime.now().strftime("%Y%m%d-%H%M%S")


def sheet_name_for(row: Mapping[str, Any]) -> str:
    return str(row.get("__sheet_name") or "Sheet1").strip() or "Sheet1"


def group_rows(rows: Iterable[Mapping[str, Any]]) -> OrderedDict[str, list[Mapping[str, Any]]]:
    grouped: OrderedDict[str, list[Mapping[str, Any]]] = OrderedDict()
    for row in rows:
        name = sheet_name_for(row)
        grouped.setdefault(name, []).append(row)
    return grouped


def _all_columns(rows: list[Mapping[str, Any]], configured: list[str]) -> list[str]:
    columns = list(configured)
    seen = set(columns)
    for row in rows:
        for key in row.keys():
            if key.startswith("__") or key in seen:
                continue
            columns.append(key)
            seen.add(key)
    return columns


def _stringify(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    return str(value)


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="E9EEF8")
    header_font = Font(bold=True, color="1F2937")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        values = [str(cell.value or "") for cell in column_cells[:80]]
        max_len = max([len(value) for value in values] or [8])
        width = max(10, min(42, max_len + 2))
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_workbook(rows: list[Mapping[str, Any]], out_dir: Path | str, filename: str | None = None) -> Path:
    if not rows:
        raise ValueError("No rows to export")
    out_dir = Path(out_dir).expanduser()
    out_dir.mkdir(parents=True, exist_ok=True)
    filename = filename or f"唯品会爆款策略追踪报表_{timestamp()}.xlsx"
    target = out_dir / filename
    if target.exists():
        stem, suffix = target.stem, target.suffix
        index = 1
        while target.exists():
            target = out_dir / f"{stem}-{index}{suffix}"
            index += 1

    grouped = group_rows(rows)
    ordered_names = [name for name in SHEET_ORDER if name in grouped]
    ordered_names.extend(name for name in grouped.keys() if name not in ordered_names)

    wb = Workbook()
    for index, sheet_name in enumerate(ordered_names):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = sheet_name[:31]
        sheet_rows = grouped[sheet_name]
        columns = _all_columns(sheet_rows, SHEET_COLUMNS.get(sheet_name, []))
        ws.append(columns)
        for row in sheet_rows:
            ws.append([_stringify(row.get(column, "")) for column in columns])
        _style_sheet(ws)

    wb.save(target)
    return target
